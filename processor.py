import io, os, json, re
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

MAIN_REQUIRED={'Account Description','Account #','State','LegEnt','Beginning Balance','Lag','JEs & Payments','Add.Entity Activity','Extract','Lag Ledger Activity','Ledger Activity','AE Tracking'}
SMW_REQUIRED={'Tax Code','Taxable Amount','Tax Amount','State','Legal Entity','GL Account'}
ID_REQUIRED={'Legal Entity','State','Net Sales','Sales Tax','Use Tax','GL Account'}

def _rows(ws):
    it=ws.iter_rows(values_only=True); hdr=[str(x).strip() if x is not None else '' for x in next(it)]
    return hdr,[dict(zip(hdr,r)) for r in it if any(v is not None for v in r)]

def _num(v):
    if v in (None,''): return 0.0
    try: return float(v)
    except: return 0.0

def validate_sheet(ws, required):
    hdr,_=_rows(ws); missing=required-set(hdr)
    if missing: raise ValueError(f'{ws.title}: missing required columns: {sorted(missing)}')

def parse_excel(main_bytes,id_bytes):
    m=load_workbook(io.BytesIO(main_bytes),data_only=True)
    i=load_workbook(io.BytesIO(id_bytes),data_only=True)
    # all tabs are read and retained in raw map; required business tabs are validated.
    raw={}
    for ws in m.worksheets: raw[f'main::{ws.title}']=_rows(ws)
    for ws in i.worksheets: raw[f'id::{ws.title}']=_rows(ws)
    if 'Blackline Monthly Balance' not in m.sheetnames or 'SMW Data' not in m.sheetnames or 'Data' not in i.sheetnames:
        raise ValueError('Required tabs: Blackline Monthly Balance, SMW Data, Data')
    validate_sheet(m['Blackline Monthly Balance'],MAIN_REQUIRED)
    validate_sheet(m['SMW Data'],SMW_REQUIRED)
    validate_sheet(i['Data'],ID_REQUIRED)
    return raw

def reconcile(raw,pdf_rows=None,user_inputs=None):
    _,black=raw['main::Blackline Monthly Balance']; _,smw=raw['main::SMW Data']; _,idr=raw['id::Data']
    pdf_rows=pdf_rows or []; user_inputs=user_inputs or {}
    state=(user_inputs.get('state') or 'ID').upper(); period=user_inputs.get('period') or 'July 2026'
    main=[]
    for r in black:
        if str(r.get('State','')).upper()!=state: continue
        x=dict(r)
        x['Open Item']=round(_num(x.get('Beginning Balance'))+_num(x.get('Lag'))+_num(x.get('JEs & Payments')),2)
        x['Actual Open Items']=round(x['Open Item']+_num(x.get('Add.Entity Activity')),2)
        x['Ending Balance']=round(x['Actual Open Items']+_num(x.get('Extract'))+_num(x.get('Lag Ledger Activity'))+_num(x.get('Ledger Activity'))+_num(x.get('AE Tracking')),2)
        main.append(x)
    entities=sorted({str(r.get('LegEnt')) for r in main if r.get('LegEnt')})
    details={}
    for ent in entities:
        acct=[r for r in main if str(r.get('LegEnt'))==ent]
        sm=[r for r in smw if str(r.get('Legal Entity'))==ent and str(r.get('State')).upper()==state]
        ids=[r for r in idr if str(r.get('Legal Entity'))==ent and str(r.get('State')).upper()==state]
        pdf=[r for r in pdf_rows if str(r.get('Legal Entity',''))==ent]
        details[ent]={
          'period':period,'state':state,'accounts':acct,
          'metrics':{
            'Beginning Balance':round(sum(_num(r.get('Beginning Balance')) for r in acct),2),
            'Prior Month Lag':round(sum(_num(r.get('Lag')) for r in acct),2),
            'JEs & Payments':round(sum(_num(r.get('JEs & Payments')) for r in acct),2),
            'Add. Entity Activity':round(sum(_num(r.get('Add.Entity Activity')) for r in acct),2),
            'Open Items':round(sum(_num(r.get('Actual Open Items')) for r in acct),2),
            'Extract SAP':round(sum(_num(r.get('Extract')) for r in acct),2),
            'Current Lag':round(sum(_num(r.get('Lag Ledger Activity')) for r in acct),2),
            'Ledger Activity':round(sum(_num(r.get('Ledger Activity')) for r in acct),2),
            'AE Tracking':round(sum(_num(r.get('AE Tracking')) for r in acct),2),
            'Ending Balance':round(sum(_num(r.get('Ending Balance')) for r in acct),2),
            'SMW Taxable Sales':round(sum(_num(r.get('Taxable Amount')) for r in sm),2),
            'SMW Tax':round(sum(_num(r.get('Tax Amount')) for r in sm),2),
            'Extract Net Sales':round(sum(_num(r.get('Net Sales')) for r in ids),2),
            'Extract Sales Tax':round(sum(_num(r.get('Sales Tax')) for r in ids),2),
            'Extract Use Tax':round(sum(_num(r.get('Use Tax')) for r in ids),2),
            'PDF Liability':round(sum(_num(r.get('Total Liability')) for r in pdf),2),
          }
        }
    return main,details

def build_output(main,details):
    wb=Workbook(); ws=wb.active; ws.title='Main Account Recon'
    cols=['Account Description','Account #','State','LegEnt','Beginning Balance','Lag','JEs & Payments','Open Item','Add.Entity Activity','Actual Open Items','Extract','Lag Ledger Activity','Ledger Activity','AE Tracking','Ending Balance']
    ws.append(cols)
    for r in main: ws.append([r.get(c) for c in cols])
    for ent in ['ABC','XYZ','LMN']:
        sh=wb.create_sheet(ent); d=details.get(ent,{'period':'','state':'','accounts':[],'metrics':{}}); m=d['metrics']
        sh['A1']='FINAL'; sh['A2']=d.get('state',''); sh['A3']=d.get('period',''); sh['A5']='RECONCILIATION SUMMARY'
        sh.append([])
        sh.append(['Metric','Calculated Value'])
        for k,v in m.items(): sh.append([k,v])
        sh.append([]); sh.append(['ACCOUNT DETAIL'])
        sh.append(cols)
        for r in d.get('accounts',[]): sh.append([r.get(c) for c in cols])
        sh.append([]); sh.append(['CONTROL CHECKS','Value'])
        sh.append(['Open Item formula check',round(sum(_num(r.get('Beginning Balance'))+_num(r.get('Lag'))+_num(r.get('JEs & Payments'))-_num(r.get('Open Item')) for r in d.get('accounts',[])),2)])
        sh.append(['Ending Balance formula check',round(sum(_num(r.get('Actual Open Items'))+_num(r.get('Extract'))+_num(r.get('Lag Ledger Activity'))+_num(r.get('Ledger Activity'))+_num(r.get('AE Tracking'))-_num(r.get('Ending Balance')) for r in d.get('accounts',[])),2)])
    thin=Side(style='thin',color='D9E2F3')
    for sh in wb.worksheets:
        sh.freeze_panes='A2'
        for cell in sh[1]: cell.fill=PatternFill('solid',fgColor='1F4E78'); cell.font=Font(color='FFFFFF',bold=True)
        for row in sh.iter_rows():
            for c in row: c.border=Border(bottom=thin); c.alignment=Alignment(vertical='top')
        for col in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(col)].width=min(38,max(12, max((len(str(sh.cell(r,col).value or '')) for r in range(1,min(sh.max_row,60)+1)),default=12)+2))
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()
