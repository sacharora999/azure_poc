import sys
from pathlib import Path
from datetime import datetime

import pandas as pd
from PySide6.QtCore import Qt, QAbstractTableModel, QThread, Signal
from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QFileDialog, QListWidget, QComboBox,
    QTableView, QMessageBox, QProgressBar, QGroupBox, QFormLayout,
    QAbstractItemView
)
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class DataFrameModel(QAbstractTableModel):
    def __init__(self, frame=None):
        super().__init__()
        self.frame = frame if frame is not None else pd.DataFrame()

    def rowCount(self, parent=None):
        return len(self.frame)

    def columnCount(self, parent=None):
        return len(self.frame.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if index.isValid() and role == Qt.ItemDataRole.DisplayRole:
            value = self.frame.iat[index.row(), index.column()]
            return "" if pd.isna(value) else str(value)
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole:
            return str(self.frame.columns[section]) if orientation == Qt.Orientation.Horizontal else str(section + 1)
        return None


class ExcelWorker(QThread):
    finished_ok = Signal(object, object, str)
    failed = Signal(str)

    def __init__(self, paths, group_col, value_col, agg, output_path):
        super().__init__()
        self.paths, self.group_col, self.value_col = paths, group_col, value_col
        self.agg, self.output_path = agg, output_path

    def run(self):
        try:
            frames = []
            for path in self.paths:
                frame = pd.read_excel(path, sheet_name=0, engine="openpyxl")
                frame.columns = [str(col).strip() for col in frame.columns]
                if frame.columns.duplicated().any():
                    raise ValueError(f"Duplicate column names in {Path(path).name}")
                frame["Source_File"] = Path(path).name
                frames.append(frame)
            data = pd.concat(frames, ignore_index=True, sort=False)
            if self.group_col not in data or self.value_col not in data:
                raise ValueError("Selected columns are missing from one or more inputs.")
            if self.agg != "count":
                numeric = pd.to_numeric(data[self.value_col], errors="coerce")
                if numeric.notna().sum() == 0:
                    raise ValueError("The value column has no numeric values for this aggregation.")
                data[self.value_col] = numeric
            pivot = pd.pivot_table(
                data, index=[self.group_col], values=[self.value_col],
                aggfunc=self.agg, fill_value=0, dropna=False
            ).reset_index()
            summary = pd.DataFrame([
                {"Metric": "Input files", "Value": len(self.paths)},
                {"Metric": "Input rows", "Value": len(data)},
                {"Metric": "Pivot rows", "Value": len(pivot)},
                {"Metric": "Grouping column", "Value": self.group_col},
                {"Metric": "Value column", "Value": self.value_col},
                {"Metric": "Aggregation", "Value": self.agg},
                {"Metric": "Generated at", "Value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            ])
            with pd.ExcelWriter(self.output_path, engine="openpyxl") as writer:
                data.to_excel(writer, sheet_name="Consolidated", index=False)
                pivot.to_excel(writer, sheet_name="Pivot_Summary", index=False)
                summary.to_excel(writer, sheet_name="Run_Summary", index=False)
                for ws in writer.book.worksheets:
                    ws.freeze_panes = "A2"
                    ws.auto_filter.ref = ws.dimensions
                    for cell in ws[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="174A67")
                        cell.alignment = Alignment(wrap_text=True)
                    for col in ws.columns:
                        letter = get_column_letter(col[0].column)
                        max_len = max((len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200]), default=10)
                        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 42)
            self.finished_ok.emit(data.head(200).copy(), pivot.head(200).copy(), self.output_path)
        except Exception as exc:
            self.failed.emit(str(exc))


class ExcelApp(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Excel Report Generator - PySide6")
        self.resize(1050, 720)
        self.paths = []
        self.worker = None
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        title = QLabel("Monthly Excel Report Generator")
        title.setStyleSheet("font-size: 22px; font-weight: bold; margin: 8px 0;")
        layout.addWidget(title)

        file_box = QGroupBox("1. Import Excel files (first worksheet of each file)")
        file_layout = QVBoxLayout(file_box)
        self.file_list = QListWidget()
        self.file_list.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        file_layout.addWidget(self.file_list)
        buttons = QHBoxLayout()
        add = QPushButton("Add Excel Files")
        add.clicked.connect(self.add_files)
        remove = QPushButton("Remove Selected")
        remove.clicked.connect(self.remove_selected)
        buttons.addWidget(add)
        buttons.addWidget(remove)
        file_layout.addLayout(buttons)
        layout.addWidget(file_box)

        config_box = QGroupBox("2. Pivot configuration")
        form = QFormLayout(config_box)
        self.group = QComboBox()
        self.value = QComboBox()
        self.agg = QComboBox()
        self.agg.addItems(["sum", "count", "mean", "min", "max"])
        form.addRow("Group by column:", self.group)
        form.addRow("Value column:", self.value)
        form.addRow("Aggregation:", self.agg)
        layout.addWidget(config_box)

        self.generate = QPushButton("Generate Excel Report")
        self.generate.clicked.connect(self.generate_report)
        layout.addWidget(self.generate)
        self.progress = QProgressBar()
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        layout.addWidget(self.progress)
        self.status = QLabel("Ready. Add Excel files to begin.")
        layout.addWidget(self.status)
        layout.addWidget(QLabel("Consolidated preview (up to 200 rows):"))
        self.data_table = QTableView()
        layout.addWidget(self.data_table, 1)
        layout.addWidget(QLabel("Pivot preview (up to 200 rows):"))
        self.pivot_table = QTableView()
        layout.addWidget(self.pivot_table, 1)

    def add_files(self):
        paths, _ = QFileDialog.getOpenFileNames(self, "Choose Excel files", "", "Excel Files (*.xlsx)")
        for path in paths:
            if path not in self.paths:
                self.paths.append(path)
                self.file_list.addItem(path)
        self.load_columns()

    def remove_selected(self):
        for item in self.file_list.selectedItems():
            path = item.text()
            self.paths.remove(path)
            self.file_list.takeItem(self.file_list.row(item))
        self.load_columns()

    def load_columns(self):
        self.group.clear()
        self.value.clear()
        if not self.paths:
            return
        try:
            cols = pd.read_excel(self.paths[0], sheet_name=0, nrows=0, engine="openpyxl").columns
            names = [str(col).strip() for col in cols]
            names.append("Source_File")
            self.group.addItems(names)
            self.value.addItems(names)
            if len(names) > 1:
                self.value.setCurrentIndex(1)
            self.status.setText(f"{len(self.paths)} input file(s) selected. Choose pivot fields.")
        except Exception as exc:
            QMessageBox.critical(self, "Invalid Excel file", str(exc))

    def generate_report(self):
        if not self.paths:
            QMessageBox.warning(self, "Missing files", "Please add at least one Excel file.")
            return
        output, _ = QFileDialog.getSaveFileName(self, "Save Excel report", "Monthly_Excel_Report.xlsx", "Excel Files (*.xlsx)")
        if not output:
            return
        if not output.lower().endswith(".xlsx"):
            output += ".xlsx"
        self.generate.setEnabled(False)
        self.progress.setRange(0, 0)
        self.status.setText("Processing Excel files...")
        self.worker = ExcelWorker(self.paths[:], self.group.currentText(), self.value.currentText(), self.agg.currentText(), output)
        self.worker.finished_ok.connect(self.done)
        self.worker.failed.connect(self.error)
        self.worker.start()

    def done(self, data_preview, pivot_preview, output):
        self.data_table.setModel(DataFrameModel(data_preview))
        self.pivot_table.setModel(DataFrameModel(pivot_preview))
        self.progress.setRange(0, 1)
        self.progress.setValue(1)
        self.generate.setEnabled(True)
        self.status.setText(f"Completed: {output}")
        QMessageBox.information(self, "Report generated", f"Saved successfully:\n{output}")

    def error(self, message):
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.generate.setEnabled(True)
        self.status.setText("Processing failed")
        QMessageBox.critical(self, "Processing error", message)


if __name__ == "__main__":
    app = QApplication(sys.argv)
    window = ExcelApp()
    window.show()
    sys.exit(app.exec())
